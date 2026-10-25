"""Try the SWP API calls on their own and write what comes back as a field sheet (Excel).

Calls made (the numbering used in the sheet):
  2  trip search   POST /api/trip/external/ExternalTrip/tripsearch/programId/{program}              -> "record"
  3  get trip      GET  /api/trip/external/ExternalTrip/programId/{program}/tripId/{trip}            -> "basic"
  4  trip detail   GET  /api/trip/external/ExternalTripDetail/programId/{program}/tripId/{trip}      -> "detail"

Run it from the backend folder (credentials are read from backend/.env, see .env.example):

  python scripts/swp_try_api.py --program 651 --limit 3                  # calls the API for 3 Arrived trips
  python scripts/swp_try_api.py --program 651 --status InTransit --limit 5
  python scripts/swp_try_api.py --from-store data/swp_cache/651          # no API: reads the trips the app stored

It writes into --out (default swp_try_output/):
  raw/<trip guid>.json     each trip's search record, get-trip and trip-detail responses, exactly as returned
  SWP_API_Fields.xlsx      one sheet per call listing every field, its type and how many trips fill it, and
                           a sheet mapping each aggregate column to the call and field it comes from
"""
from __future__ import annotations

import argparse
import collections
import glob
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # so `app` imports when run from anywhere
sys.path.insert(0, str(Path(__file__).resolve().parent))      # and swp_columns, next to this script

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

# ---- which aggregate column(s) each field feeds (the "Aggregate column(s)" and "How" columns of the field sheets) ------
# path -> (aggregate columns, how). How: Direct, Converted (time zone / unit), Computed (worked out from several
# fields or the readings), Fallback (used only when the first source is empty), Available (the field is there and
# the column exists, but the app does not fill it yet), App (not an export column).
_TIME_COLS = ["Time Below Low Hours", "Time Below Ideal Hours", "Time Above Ideal Hours", "Time Above High Hours"]
_MINUTE_COLS = ["Value Minutes Below Low", "Value Minutes Below Ideal", "Value Minutes Above Ideal", "Value Minutes Above High"]
_STATS = ["Mean Value", "Standard Deviation", "Min Value", "Max Value"]
_AVAILABLE = "Available (empty in the export, not filled yet)"
AGG = {
    "basic": {
        "serialNumbers[]": (["Serial Number"], "Fallback (only for a trip with no readings; the detail has the same serials)"),
        "quality": (["Quality Status"], "Direct"),
        "qualityComment": (["Quality Comment"], "Direct"),
        "parentSupplierId": (["Parent Supplier"], "Direct"),
        "internalTripId": (["Internal Trip ID"], "Direct"),
        "notes": (["Trip Note"], "Direct"),
    },
    "detail": {
        "id": (["Trip ID"], "Direct"),
        "tripID": (["SWP Trip Key"], "App (the trip GUID; not an export column)"),
        "createDate": (["Trip Create Date CET", "Shipment Created Date/Time"], "Converted (UTC to CET)"),
        "tripStatusGuid": (["Trip Status"], "Fallback (the search record gives the text)"),
        "alarmed": (["Is Alarmed"], "Converted (true/false to Yes/No)"),
        "trailerID": (["Trailer ID"], "Direct"),
        "notes": (["Trip Note"], "Fallback"),
        "tripActualArrivalTime": (["Actual Arrival Time CET", "Segment End Date Time", "Segment Length (Days)"], "Converted / Computed"),
        "tripOrigin.actualDepartureTime": (["Actual Departure Time CET", "Segment Start Date Time", "Segment Length (Days)"], "Converted / Computed"),
        "tripOrigin.plannedDepartureTime": (["Planned Departure Time CET"], "Converted (UTC to CET)"),
        "tripOrigin.carrier.name": (["Carrier"], "Direct"),
        "tripOrigin.managingCarrier": (["Managing Carrier"], "Direct (null for every trip, so the name, managingCarrier.name, is empty)"),
        "tripOrigin.driverName": (["Driver Name"], "Direct"),
        "tripOrigin.driverCell": (["Driver Cell"], "Direct"),
        "tripOrigin.internalTripID": (["Internal Trip ID"], "Fallback"),
        "tripOrigin.locations.locationName": (["Origin"], "Direct"),
        "destinations[].locationName": (["Destination"], "Direct"),
        "destinations[].orderNumber": (["Order Number"], "Direct"),
        "destinations[].poNumber": (["PO Number"], "Direct"),
        "destinations[].modeOfTransportation": (["Mode Of Transportation"], "Direct"),
        "destinations[].plannedArrivalTime": (["Planned Arrival Time CET"], "Converted (UTC to CET)"),
        "destinations[].actualArrivalTime": (["Actual Arrival Time CET"], "Fallback"),
        "destinations[].products[].productName": (["Product"], "Direct"),
        "containers[].containerName": (["Container Name"], _AVAILABLE),
        "containers[].zones[].zoneName": (["Zone Name"], _AVAILABLE),
        "segments[].segmentName": (["Segment Name"], "Available (the app writes 'Full Trip' instead)"),
        "segments[].segmentType": (["Segment Type"], _AVAILABLE),
        "segments[].segmentNumber": (["Segment Number"], _AVAILABLE),
        "segments[].startTime": (["Segment Start Date Time"], "Converted (UTC to CET)"),
        "segments[].endTime": (["Segment End Date Time"], "Converted (UTC to CET)"),
        "containers[].zones[].zoneProducts[].minTemperature": (["Limit Low", "Time Below Low Hours", "Value Minutes Below Low"], "Converted (F to C) / Computed"),
        "containers[].zones[].zoneProducts[].idealTemperature": (
            ["Limit Ideal", "Time Below Ideal Hours", "Time Above Ideal Hours", "Value Minutes Below Ideal", "Value Minutes Above Ideal"],
            "Converted (F to C) / Computed",
        ),
        "containers[].zones[].zoneProducts[].maxTemperature": (["Limit High", "Time Above High Hours", "Value Minutes Above High"], "Converted (F to C) / Computed"),
        "sensorData[].serialNumber": (["Serial Number"], "Direct (one row per logger)"),
        "sensorData[].temperatureData[]": (["Sensor Type", "Degree Unit"], "Direct (a Temperature row; unit C)"),
        "sensorData[].temperatureData[].value": (_STATS + ["MKT Statistics Value"] + _TIME_COLS + _MINUTE_COLS, "Computed (F to C, then statistics)"),
        "sensorData[].temperatureData[].dateTime": (_TIME_COLS + _MINUTE_COLS, "Computed (how long each reading counts)"),
        "sensorData[].lightData[]": (["Sensor Type", "Degree Unit"], "Direct (a Light row; unit %)"),
        "sensorData[].lightData[].value": (_STATS, "Computed (statistics of the readings)"),
    },
    "record": {
        "tripStatus": (["Trip Status"], "Direct (text; also the In Transit / Arrived split)"),
        "tripStatusGuid": (["Trip Status"], "Fallback"),
        "createdBy": (["Shipment Created By"], "Direct"),
        "tripStatistics.mkt": (["MKT Statistics Value"], "Converted (F to C); computed from the readings when empty"),
        "tripStatistics.tempSpec.low": (["Limit Low"], "Fallback (F to C)"),
        "tripStatistics.tempSpec.high": (["Limit High"], "Fallback (F to C)"),
        "tripStatistics.tempVariance.mean": (["Mean Value"], "Fallback (a trip with no readings; F to C)"),
        "tripStatistics.tempVariance.sd": (["Standard Deviation"], "Fallback (a trip with no readings)"),
        "tripStatistics.tempReading.min": (["Min Value"], "Fallback (a trip with no readings; F to C)"),
        "tripStatistics.tempReading.max": (["Max Value"], "Fallback (a trip with no readings; F to C)"),
        "lastUpdatedDate": ([], "Not a column: tells whether a stored trip is still current"),
        "tripID": ([], "Not a column: the trip GUID"),
    },
}
SHEETS = [
    ("Get trip (3)", "basic", "GET /api/trip/external/ExternalTrip/programId/{program}/tripId/{trip}"),
    ("Trip detail (4)", "detail", "GET /api/trip/external/ExternalTripDetail/programId/{program}/tripId/{trip}"),
    ("Search record (2)", "record", "POST /api/trip/external/ExternalTrip/tripsearch/programId/{program}"),
]


# ---- reading the responses ---------------------------------------------------------------------------------------

def walk(obj, path: str, out: dict) -> None:
    """Every field path of a response with its type; a list is written `name[]` and its first item described."""
    if isinstance(obj, dict):
        for key, value in obj.items():
            walk(value, f"{path}.{key}" if path else key, out)
    elif isinstance(obj, list):
        if not obj:
            out[path + "[]"] = "empty"
            return
        out[path + "[]"] = "list"
        walk(obj[0], path + "[]", out)
    else:
        out[path] = "null" if obj is None else type(obj).__name__


def field_table(payloads: list[dict], part: str) -> tuple[int, list[list]]:
    """(trips counted, rows [field, type, filled, of, aggregate columns, how]) for one part of the payloads."""
    stats: dict[str, list] = collections.defaultdict(lambda: [0, set()])
    counted = 0
    for payload in payloads:
        data = payload.get(part) or {}
        if not data:
            continue  # e.g. trips stored before the search record was kept
        counted += 1
        found: dict[str, str] = {}
        walk(data, "", found)
        for path, kind in found.items():
            stats[path][1].add(kind)
            if kind not in ("null", "empty"):
                stats[path][0] += 1
    rows = []
    for path in sorted(stats):
        filled, kinds = stats[path]
        kind = "/".join(sorted(k for k in kinds if k not in ("null", "empty")) or sorted(kinds))
        columns, how = AGG[part].get(path, ([], ""))
        rows.append([path, kind, filled, counted, ", ".join(columns), how])
    return counted, rows


# ---- the workbook ------------------------------------------------------------------------------------------------

def style(sheet, widths: list[int], header_row: int) -> None:
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="010198")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(widths))}{sheet.max_row}"


def write_workbook(payloads: list[dict], path: Path, program: str) -> dict[str, int]:
    wb = Workbook()
    info = wb.active
    info.title = "Read me"
    info.append(["SWP API fields"])
    info["A1"].font = Font(bold=True, size=16)
    info.append([f"Program {program}. Built from {len(payloads)} trips, from the responses exactly as the API returned them."])
    info.append(["Each of the other sheets lists every field of one call, with its type and how many trips have a value for it."])
    info.append(["Green rows feed the aggregate file: 'Aggregate column(s)' names the column(s) of Tabular Shipment Data, 'How' says how."])
    info.append(["'Search (2) vs Detail (4)' lists every aggregate column with the field in each call, whether it is in both or only one, and whether the two values agree on the arrived trips."])
    info.append([])
    info.append(["Call", "Endpoint"])
    for title, _part, endpoint in SHEETS:
        info.append([title, endpoint])
    info.append(["Token", "POST /api/security-external/external/auth/connect/token"])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 100

    from swp_columns import matrix  # next to this script

    sources = wb.create_sheet("Search (2) vs Detail (4)")
    sources.append(["Aggregate column", "Trip search (2): field", "Trip detail (4): field (raw data)", "Available in", "Values agree (arrived trips)", "Note"])
    fills = {"Both": "E3F4E6", "Search (2) only": "E4EFFD", "Detail (4) only": "FFF3CD", "Get trip (3) only": "EEEEEE", "Neither": "EEEEEE"}
    for row in matrix(payloads):
        sources.append(row)
        sources.cell(row=sources.max_row, column=4).fill = PatternFill("solid", fgColor=fills.get(row[3], "FFFFFF"))
    style(sources, [32, 58, 58, 18, 34, 70], 1)

    counts: dict[str, int] = {}
    for title, part, endpoint in SHEETS:
        counted, rows = field_table(payloads, part)
        counts[title] = len(rows)
        sheet = wb.create_sheet(title)
        sheet.append([f"{endpoint}  -  {len(rows)} fields over {counted} trips"])
        sheet["A1"].font = Font(bold=True)
        sheet.append(["Field", "Type", "Trips with a value", "Trips counted", "Aggregate column(s)", "How"])
        for row in rows:
            sheet.append(row)
            if row[4]:  # a field that feeds an aggregate column is marked green
                for cell in sheet[sheet.max_row]:
                    cell.fill = PatternFill("solid", fgColor="E3F4E6")
        style(sheet, [64, 10, 16, 14, 64, 62], 2)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return counts


# ---- getting the trips -------------------------------------------------------------------------------------------

def from_store(folder: str) -> list[dict]:
    files = sorted(glob.glob(str(Path(folder) / "trips" / "*.json")))
    if not files:
        sys.exit(f"No stored trips in {folder}/trips")
    payloads = []
    for f in files:
        payload = json.loads(Path(f).read_text(encoding="utf-8"))
        payload.setdefault("tripID", Path(f).stem)  # trips stored by an older version carry no id inside
        payloads.append(payload)
    return payloads


def from_api(program: str, status: str, limit: int) -> list[dict]:
    import time

    import requests

    from app import config
    from app.services.swp import client

    program = program or config.SWP_PROGRAM_ID
    missing = client.missing_settings(program)
    if missing:
        sys.exit("Set in backend/.env: " + ", ".join(missing))
    token = client.get_token()
    print("token ok")

    def call(method: str, path: str, **kwargs) -> dict:
        # a plain call with the first trip key, paced at one call every 3 seconds (this script makes only a handful)
        response = requests.request(
            method, config.SWP_BASE_URL + path, timeout=120, **kwargs,
            headers={"Ocp-Apim-Subscription-Key": config.SWP_TRIP_KEY, "Authorization": f"Bearer {token}"},
        )
        time.sleep(config.SWP_CALL_GAP_SECONDS)
        response.raise_for_status()
        return response.json()

    page = call("POST", f"/api/trip/external/ExternalTrip/tripsearch/programId/{program}", json={"keyword": status, "from": 0})
    records = (page.get("records") or [])[:limit]
    print(f"search '{status}': {page.get('overallTotal')} trips in all, taking {len(records)}")
    payloads = []
    for record in records:
        trip_id = record["tripID"]
        basic = call("GET", f"/api/trip/external/ExternalTrip/programId/{program}/tripId/{trip_id}")
        detail = call("GET", f"/api/trip/external/ExternalTripDetail/programId/{program}/tripId/{trip_id}")
        print(f"  {trip_id}: get trip + trip detail")
        payloads.append({"tripID": trip_id, "record": record, "basic": basic, "detail": detail})
    return payloads


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--program", default="", help="program id (default: SWP_PROGRAM_ID)")
    parser.add_argument("--status", default="Arrived", help="status to search for: Arrived, InTransit, ... (default Arrived)")
    parser.add_argument("--limit", type=int, default=3, help="how many trips to fetch from the API (default 3)")
    parser.add_argument("--from-store", metavar="FOLDER", help="skip the API and read the trips the app stored, e.g. data/swp_cache/651")
    parser.add_argument("--out", default="swp_try_output", help="folder to write into")
    args = parser.parse_args()

    payloads = from_store(args.from_store) if args.from_store else from_api(args.program, args.status, args.limit)
    out = Path(args.out)
    (out / "raw").mkdir(parents=True, exist_ok=True)
    for payload in payloads:
        (out / "raw" / f"{payload['tripID']}.json").write_text(json.dumps(payload, indent=1), encoding="utf-8")
    counts = write_workbook(payloads, out / "SWP_API_Fields.xlsx", args.program or (args.from_store or ""))
    print(f"{len(payloads)} trips -> {out}/SWP_API_Fields.xlsx", counts)


if __name__ == "__main__":
    main()
