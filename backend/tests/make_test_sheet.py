"""Builds docs/PM-AI_Test_Cases.xlsx from the tests themselves, so the sheet and the Python files never drift apart.

Every test whose docstring reads  "TC-<AREA>-<NNN> | scenario | expected result"  becomes one row. The tests are run and
the result (Pass / Fail / Error) is written next to each row.

    python -m tests.make_test_sheet            (from the backend folder)
"""
from __future__ import annotations

import datetime
import re
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).resolve().parents[2] / "docs" / "PM-AI_Test_Cases.xlsx"
AREAS = {"SWP": "SWP integration (SensiWatch API)"}
ROW = re.compile(r"^\s*(TC-([A-Z]+)-\d+)\s*\|\s*(.+?)\s*\|\s*(.+?)\s*$")


def cases(suite):
    for item in suite:
        if isinstance(item, unittest.TestSuite):
            yield from cases(item)
        else:
            doc = (item.shortDescription() or "").strip()
            m = ROW.match(doc)
            if m:
                yield item, m.groups()


def main() -> None:
    suite = unittest.defaultTestLoader.discover("tests", top_level_dir=".")
    found = list(cases(suite))
    result = unittest.TestResult()
    for test, _ in found:
        test.run(result)
    failed = {t.id() for t, _ in result.failures}
    errored = {t.id() for t, _ in result.errors}

    wb = Workbook()
    ws = wb.active
    ws.title = "Test cases"
    ws.append(["Test case ID", "Area", "Scenario (given / when)", "Expected result", "Python test", "Last result", "Run on"])
    today = datetime.date.today().isoformat()
    for test, (tc, area, scenario, expected) in sorted(found, key=lambda f: f[1][0]):
        status = "Fail" if test.id() in failed else "Error" if test.id() in errored else "Pass"
        file = test.id().split(".")[1] + ".py"
        ws.append([tc, AREAS.get(area, area), scenario, expected, f"tests/{file} :: {test.id().split('.', 2)[2]}", status, today])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E9E")
    for col, width in zip("ABCDEFG", (14, 30, 62, 62, 60, 12, 12)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        row[5].fill = PatternFill("solid", fgColor="C6EFCE" if row[5].value == "Pass" else "FFC7CE")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"{len(found)} test cases, {len(failed)} failed, {len(errored)} errors -> {OUT}")


if __name__ == "__main__":
    main()
