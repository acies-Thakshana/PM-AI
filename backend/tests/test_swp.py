"""Unit tests of the SWP (SensiWatch) integration: conversions, the key pool, the store and the tables built from trips.

Every test's docstring is one row of docs/PM-AI_Test_Cases.xlsx (see tests/make_test_sheet.py):

    TC-SWP-001 | Scenario in plain words | Expected result

Run:  python -m unittest tests.test_swp -v          (from the backend folder)
"""
import io
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import pandas as pd
from openpyxl import load_workbook

from app import config
from app.services.swp import cache, client, preprocess, transform


def record(**over) -> dict:
    """A trip-search record as SWP returns it (temperatures in Fahrenheit)."""
    base = {
        "tripID": "00000001-0000-4000-8000-000000000000", "id": 4700001, "tripStatus": "Arrived",
        "createDate": "2026-10-01T01:00:00Z", "lastUpdatedDate": "2026-10-05T10:00:00Z", "tripLength": "3.12:00:00",
        "tripActualArrivalTime": "2026-10-05T10:00:00+00:00",
        "tripOrigin": {"actualDepartureTime": "2026-10-01T22:00:00+00:00", "carrier": {"name": "Acme"}, "locations": [{"locationName": "Hamburg"}]},
        "destinations": [{"locationName": "Berlin", "products": [{"productName": "Figs"}]}],
        "tripStatistics": {"tempVariance": {"mean": 41.0, "sd": 1.8}},
    }
    base.update(over)
    return base


class ConversionTest(unittest.TestCase):
    def test_fahrenheit_to_celsius(self):
        """TC-SWP-001 | A temperature of 41 °F comes from SWP | It is shown as 5.0 °C"""
        with mock.patch.object(config, "SWP_TEMPERATURE_FAHRENHEIT", True):
            self.assertAlmostEqual(transform.f_to_c(41), 5.0)

    def test_celsius_kept_when_conversion_is_off(self):
        """TC-SWP-002 | Conversion is switched off (SWP_TEMPERATURE_FAHRENHEIT=0) | The value is left as it is"""
        with mock.patch.object(config, "SWP_TEMPERATURE_FAHRENHEIT", False):
            self.assertEqual(transform.f_to_c(41), 41.0)

    def test_missing_temperature(self):
        """TC-SWP-003 | The temperature is missing | The result is empty, not an error"""
        self.assertIsNone(transform.f_to_c(None))
        self.assertIsNone(transform.f_to_c(float("nan")))

    def test_trip_length_in_days(self):
        """TC-SWP-004 | The trip length is 3 days 12 hours (3.12:00:00) | Segment days are 3.5"""
        self.assertEqual(transform.duration_days("3.12:00:00"), 3.5)

    def test_trip_length_under_a_day(self):
        """TC-SWP-005 | The trip length has no day part (12:00:00) | Segment days are 0.5"""
        self.assertEqual(transform.duration_days("12:00:00"), 0.5)

    def test_trip_length_unreadable(self):
        """TC-SWP-006 | The trip length is empty or not a duration | The result is empty"""
        self.assertIsNone(transform.duration_days(None))
        self.assertIsNone(transform.duration_days("n/a"))

    def test_cet_in_summer(self):
        """TC-SWP-007 | A UTC time in July is converted to CET | Two hours are added (summer time)"""
        self.assertEqual(transform.to_cet("2026-07-01T10:00:00Z"), pd.Timestamp("2026-07-01 12:00:00"))

    def test_cet_in_winter(self):
        """TC-SWP-008 | A UTC time in January is converted to CET | One hour is added"""
        self.assertEqual(transform.to_cet("2026-01-01T10:00:00Z"), pd.Timestamp("2026-01-01 11:00:00"))

    def test_pick_takes_first_item_of_a_list(self):
        """TC-SWP-009 | A value is nested inside a list (locations[0].locationName) | The first item's value is returned"""
        self.assertEqual(transform.pick({"a": [{"b": 7}, {"b": 9}]}, "a", "b"), 7)

    def test_pick_missing_path(self):
        """TC-SWP-010 | A field on the path does not exist | The result is empty, not an error"""
        self.assertIsNone(transform.pick({"a": {}}, "a", "b", "c"))


class StatusTest(unittest.TestCase):
    def test_status_from_text(self):
        """TC-SWP-011 | The record says tripStatus "In Transit" | The status is InTransit"""
        self.assertEqual(client.status_name({"tripStatus": "In Transit"}), "InTransit")

    def test_status_from_guid(self):
        """TC-SWP-012 | The record has only a status GUID of an arrived trip | The status is Arrived"""
        guid = client.STATUS_GUIDS["arrived"]
        self.assertEqual(client.status_name({"tripStatusGuid": guid}), "Arrived")

    def test_status_unknown(self):
        """TC-SWP-013 | The record has no usable status | The status is Unknown"""
        self.assertEqual(client.status_name({}), "Unknown")


class KeyPoolTest(unittest.TestCase):
    def setUp(self):
        client._pacers[:] = []
        self.keys = mock.patch.object(config, "SWP_TRIP_KEY", "k1"), mock.patch.object(config, "SWP_TRIP_KEY_2", "k2")
        self.gap = mock.patch.object(config, "SWP_CALL_GAP_SECONDS", 3.0)
        for p in (*self.keys, self.gap):
            p.start()
            self.addCleanup(p.stop)

    def test_two_keys_share_the_calls(self):
        """TC-SWP-014 | Four calls are made with two keys | The keys alternate: k1, k2, k1, k2"""
        with mock.patch.object(client.time, "sleep"):
            used = [client._reserve().key for _ in range(4)]
        self.assertEqual(used, ["k1", "k2", "k1", "k2"])

    def test_a_key_waits_the_gap_between_its_calls(self):
        """TC-SWP-015 | A key was just used | Its next call is held back by about 3 seconds"""
        waits = []
        with mock.patch.object(client.time, "sleep", side_effect=waits.append):
            for _ in range(4):
                client._reserve()
        self.assertEqual(len(waits), 2)                      # only the 3rd and 4th call have to wait
        self.assertTrue(all(2.5 < w <= 3.0 for w in waits), waits)

    def test_one_key_only(self):
        """TC-SWP-016 | Only the first key is set | All calls use it, one gap apart"""
        with mock.patch.object(config, "SWP_TRIP_KEY_2", ""), mock.patch.object(client.time, "sleep") as sleep:
            used = [client._reserve().key for _ in range(2)]
        self.assertEqual(used, ["k1", "k1"])
        self.assertEqual(sleep.call_count, 1)

    def test_rate_limited_key_rests(self):
        """TC-SWP-017 | SWP answers 429 (too many calls) on a key | That key rests for 65 seconds"""
        with mock.patch.object(client.time, "sleep"):
            pacer = client._reserve()
        client._rest(pacer)
        self.assertGreaterEqual(pacer.next_at - client.time.monotonic(), 64)


class StoreTest(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        for p in (mock.patch.object(cache, "_DIR", Path(tmp.name)), mock.patch.object(config, "SWP_CACHE", True)):
            p.start()
            self.addCleanup(p.stop)

    def test_trip_is_stored_and_read_back(self):
        """TC-SWP-018 | A trip's raw answer is saved | It is read back exactly as saved"""
        payload = {"tripID": "t1", "record": {"id": 1}, "detail": {"sensorData": [1, 2]}}
        cache.save_trip("651", "t1", payload)
        self.assertEqual(cache.load_trip("651", "t1"), payload)

    def test_unknown_trip(self):
        """TC-SWP-019 | A trip that was never saved is asked for | The answer is empty"""
        self.assertIsNone(cache.load_trip("651", "nope"))

    def test_corrupt_file_is_a_miss(self):
        """TC-SWP-020 | A stored file is damaged | It is treated as missing, not as an error"""
        cache.save_trip("651", "t1", {"a": 1})
        cache._path("651", "trips", "t1.json").write_text("{broken", encoding="utf-8")
        self.assertIsNone(cache.load_trip("651", "t1"))

    def test_store_switched_off(self):
        """TC-SWP-021 | The store is switched off (SWP_CACHE=0) | Nothing is written to disk"""
        with mock.patch.object(config, "SWP_CACHE", False):
            cache.save_trip("651", "t9", {"a": 1})
        self.assertFalse(cache._path("651", "trips", "t9.json").exists())

    def test_programs_are_kept_apart(self):
        """TC-SWP-022 | Two programs have a trip with the same id | Each program keeps its own copy"""
        cache.save_trip("651", "t1", {"p": 651})
        cache.save_trip("652", "t1", {"p": 652})
        self.assertEqual((cache.load_trip("651", "t1")["p"], cache.load_trip("652", "t1")["p"]), (651, 652))


class StoredPayloadTest(unittest.TestCase):
    def _stored(self, updated):
        return {"tripID": "t1", "record": {"lastUpdatedDate": updated}, "detail": {"sensorData": []}}

    def test_arrived_trip_is_never_fetched_again(self):
        """TC-SWP-023 | An arrived trip is stored and SWP now reports a newer update | The stored copy is still used"""
        rec = {"tripID": "t1", "tripStatus": "Arrived", "lastUpdatedDate": "2026-10-09"}
        with mock.patch.object(cache, "load_trip", return_value=self._stored("2026-10-01")):
            self.assertIn("t1", preprocess.stored_payloads("651", ["t1"], {"t1": rec}))

    def test_trip_on_its_way_changes(self):
        """TC-SWP-024 | An in-transit trip changed since it was stored | The stored copy is not used"""
        rec = {"tripID": "t1", "tripStatus": "InTransit", "lastUpdatedDate": "2026-10-09"}
        with mock.patch.object(cache, "load_trip", return_value=self._stored("2026-10-01")):
            self.assertNotIn("t1", preprocess.stored_payloads("651", ["t1"], {"t1": rec}))

    def test_trip_on_its_way_unchanged(self):
        """TC-SWP-025 | An in-transit trip has not changed since it was stored | The stored copy is used"""
        rec = {"tripID": "t1", "tripStatus": "InTransit", "lastUpdatedDate": "2026-10-01"}
        with mock.patch.object(cache, "load_trip", return_value=self._stored("2026-10-01")):
            self.assertIn("t1", preprocess.stored_payloads("651", ["t1"], {"t1": rec}))

    def test_stored_without_details(self):
        """TC-SWP-026 | A trip is stored but its detail is empty | It counts as not stored"""
        rec = {"tripID": "t1", "tripStatus": "Arrived", "lastUpdatedDate": "x"}
        with mock.patch.object(cache, "load_trip", return_value={"tripID": "t1", "record": rec, "detail": {}}):
            self.assertEqual(preprocess.stored_payloads("651", ["t1"], {"t1": rec}), {})


class SearchRecordTest(unittest.TestCase):
    def test_kpis_from_the_search_record(self):
        """TC-SWP-027 | A trip search record has mean 41 °F and length 3.12:00:00 | Mean Temp 5.0 °C and Segment Days 3.5"""
        k = preprocess.record_kpis(record())
        self.assertEqual((k["mean_temp"], k["segment_days"]), (5.0, 3.5))

    def test_names_from_the_search_record(self):
        """TC-SWP-028 | A trip search record names origin, destination, carrier and product | They are returned as text"""
        k = preprocess.record_kpis(record())
        self.assertEqual((k["origin"], k["destination"], k["carrier"], k["product"]), ("Hamburg", "Berlin", "Acme", "Figs"))

    def test_record_without_statistics(self):
        """TC-SWP-029 | A search record has no statistics and no length | Mean Temp and Segment Days are empty"""
        k = preprocess.record_kpis(record(tripStatistics=None, tripLength=None))
        self.assertEqual((k["mean_temp"], k["segment_days"]), (None, None))

    def test_table_has_the_export_columns(self):
        """TC-SWP-030 | The audit table is built from search records | It has the export's columns in the export's order"""
        table = transform.record_table([record()])
        self.assertEqual(list(table.columns), transform.EXPORT_COLUMNS)

    def test_table_row_values(self):
        """TC-SWP-031 | One search record becomes one temperature row | Origin, product, mean (°C) and segment length are filled"""
        row = transform.record_table([record()]).iloc[0]
        self.assertEqual((row["Origin"], row["Product"], row["Segment Length (Days)"]), ("Hamburg", "Figs", 3.5))
        self.assertAlmostEqual(row["Mean Value"], 5.0, places=1)

    def test_light_row_comes_first(self):
        """TC-SWP-032 | A search record has light statistics | A Light row is added before the Temperature row"""
        stats = {"tempVariance": {"mean": 41.0, "sd": 1.0}, "lightVariance": {"mean": 2.0, "sd": 1.0}, "lightReadings": {"min": 0, "max": 5}}
        rows = transform.record_table([record(tripStatistics=stats)])
        self.assertEqual(list(rows["Sensor Type"]), ["Light", "Temperature"])

    def test_sensor_filter(self):
        """TC-SWP-033 | Only Temperature is selected on the Configure page | Light rows are left out"""
        stats = {"tempVariance": {"mean": 41.0}, "lightVariance": {"mean": 2.0}}
        rows = transform.record_table([record(tripStatistics=stats)], sensors=["Temperature"])
        self.assertEqual(list(rows["Sensor Type"]), ["Temperature"])


class ExcelTest(unittest.TestCase):
    def test_excel_layout(self):
        """TC-SWP-034 | The aggregate table is written to Excel | Title in A1, the export's header texts in row 3, data from row 4"""
        table = transform.record_table([record()])
        sheet = load_workbook(io.BytesIO(preprocess.excel_bytes(table))).active
        self.assertEqual(sheet["A1"].value, "Tabular Shipment Data")
        self.assertEqual([c.value for c in sheet[3]], preprocess.EXCEL_HEADERS)
        self.assertEqual(sheet.max_row, 4)

    def test_excel_has_59_columns(self):
        """TC-SWP-035 | The aggregate Excel is written | It has exactly the 59 columns of the SensiWatch export"""
        self.assertEqual(len(preprocess.AGGREGATE_COLUMNS), 59)
        self.assertNotIn("SWP Trip Key", preprocess.AGGREGATE_COLUMNS)


if __name__ == "__main__":
    unittest.main()
