"""
Deterministic PowerPoint builder functions. This is the ONLY code in the
system that touches python-pptx / PPTX XML. Agent 2 never writes to the file
directly -- it calls these functions as tools with plain arguments (title,
bullet text, table rows, ...), and this module is entirely responsible for
layout, fonts, colors and rendering fidelity. That separation is what fixes
the "failed to emulate template / broken rendering" failure mode from the
original Copilot attempt.
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.services.report_template import BRAND, PROGRAM_TITLE

PRIMARY = RGBColor.from_string(BRAND["primary_hex"])
SECONDARY = RGBColor.from_string(BRAND["secondary_hex"])
ACCENT = RGBColor.from_string(BRAND["accent_hex"])
DARK_TEXT = RGBColor.from_string(BRAND["dark_text_hex"])
LIGHT_BG = RGBColor.from_string(BRAND["light_bg_hex"])
WHITE = RGBColor.from_string("FFFFFF")

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


class PptxBuilder:
    """Accumulates slides for one report; one instance per report-generation job."""

    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W
        self.prs.slide_height = SLIDE_H
        self._blank = self.prs.slide_layouts[6]
        self.slide_count = 0

    # -- internal styling helpers -------------------------------------------------
    def _new_slide(self):
        slide = self.prs.slides.add_slide(self._blank)
        self.slide_count += 1
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = WHITE
        return slide

    def _header(self, slide, heading: str, band_color=PRIMARY):
        band = slide.shapes.add_shape(1, 0, 0, SLIDE_W, Inches(1.0))  # MSO_SHAPE.RECTANGLE = 1
        band.fill.solid()
        band.fill.fore_color.rgb = band_color
        band.line.fill.background()
        tb = band.text_frame
        tb.margin_left, tb.margin_top = Inches(0.4), Inches(0.15)
        p = tb.paragraphs[0]
        p.text = heading
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.name = BRAND["font_heading"]
        return band

    def _footer(self, slide):
        tb = slide.shapes.add_textbox(Inches(0.4), SLIDE_H - Inches(0.4), Inches(8), Inches(0.3))
        p = tb.text_frame.paragraphs[0]
        p.text = PROGRAM_TITLE
        p.font.size = Pt(9)
        p.font.color.rgb = RGBColor.from_string("888888")
        page = slide.shapes.add_textbox(SLIDE_W - Inches(1.2), SLIDE_H - Inches(0.4), Inches(0.8), Inches(0.3))
        pp = page.text_frame.paragraphs[0]
        pp.text = str(self.slide_count)
        pp.font.size = Pt(9)
        pp.alignment = PP_ALIGN.RIGHT
        pp.font.color.rgb = RGBColor.from_string("888888")

    # -- tools (one per SLIDE_SEQUENCE entry) --------------------------------------
    def add_title_slide(self, title: str, subtitle: str, period: str):
        slide = self._new_slide()
        band = slide.shapes.add_shape(1, 0, Inches(2.6), SLIDE_W, Inches(2.3))
        band.fill.solid()
        band.fill.fore_color.rgb = PRIMARY
        band.line.fill.background()
        tb = band.text_frame
        tb.word_wrap = True
        p = tb.paragraphs[0]
        p.text = title
        p.font.size = Pt(40)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        p2 = tb.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(18)
        p2.font.color.rgb = WHITE
        p2.alignment = PP_ALIGN.CENTER

        period_box = slide.shapes.add_textbox(Inches(0), Inches(5.2), SLIDE_W, Inches(0.6))
        pp = period_box.text_frame.paragraphs[0]
        pp.text = period
        pp.font.size = Pt(16)
        pp.font.color.rgb = DARK_TEXT
        pp.alignment = PP_ALIGN.CENTER
        return {"status": "ok", "slide": "title"}

    def add_kpi_slide(self, heading: str, kpis: list[dict]):
        slide = self._new_slide()
        self._header(slide, heading)
        n = max(len(kpis), 1)
        card_w = (SLIDE_W - Inches(0.8) - Inches(0.3) * (n - 1)) / n
        x = Inches(0.4)
        for kpi in kpis:
            card = slide.shapes.add_shape(1, x, Inches(1.6), card_w, Inches(3.0))
            card.fill.solid()
            card.fill.fore_color.rgb = LIGHT_BG
            status = str(kpi.get("status", "")).lower()
            card.line.color.rgb = ACCENT if status == "at_risk" else PRIMARY
            card.line.width = Pt(2)
            tf = card.text_frame
            tf.word_wrap = True
            tf.margin_top = Inches(0.25)
            p1 = tf.paragraphs[0]
            p1.text = str(kpi.get("value", ""))
            p1.font.size = Pt(34)
            p1.font.bold = True
            p1.font.color.rgb = PRIMARY if status != "at_risk" else RGBColor.from_string("B71C1C")
            p1.alignment = PP_ALIGN.CENTER
            p2 = tf.add_paragraph()
            p2.text = str(kpi.get("label", ""))
            p2.font.size = Pt(13)
            p2.font.color.rgb = DARK_TEXT
            p2.alignment = PP_ALIGN.CENTER
            p3 = tf.add_paragraph()
            p3.text = f"Target: {kpi.get('target', 'n/a')}"
            p3.font.size = Pt(11)
            p3.font.color.rgb = RGBColor.from_string("666666")
            p3.alignment = PP_ALIGN.CENTER
            x += card_w + Inches(0.3)
        self._footer(slide)
        return {"status": "ok", "slide": "kpi", "kpi_count": len(kpis)}

    def add_narrative_slide(self, heading: str, paragraphs: list[str]):
        slide = self._new_slide()
        self._header(slide, heading, band_color=SECONDARY)
        box = slide.shapes.add_textbox(Inches(0.6), Inches(1.3), SLIDE_W - Inches(1.2), Inches(5.5))
        tf = box.text_frame
        tf.word_wrap = True
        for i, para in enumerate(paragraphs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = para
            p.font.size = Pt(16)
            p.font.color.rgb = DARK_TEXT
            p.space_after = Pt(14)
        self._footer(slide)
        return {"status": "ok", "slide": "narrative", "paragraph_count": len(paragraphs)}

    def add_table_slide(self, heading: str, columns: list[str], rows: list[list[str]]):
        slide = self._new_slide()
        self._header(slide, heading, band_color=PRIMARY)
        n_rows, n_cols = len(rows) + 1, len(columns)
        table_shape = slide.shapes.add_table(
            n_rows, n_cols, Inches(0.5), Inches(1.3), SLIDE_W - Inches(1.0), Inches(5.6)
        )
        table = table_shape.table
        for c, col_name in enumerate(columns):
            cell = table.cell(0, c)
            cell.text = str(col_name)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PRIMARY
            para = cell.text_frame.paragraphs[0]
            para.font.bold = True
            para.font.size = Pt(12)
            para.font.color.rgb = WHITE
        for r, row in enumerate(rows, start=1):
            for c, val in enumerate(row):
                cell = table.cell(r, c)
                cell.text = "" if val is None else str(val)
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_BG if r % 2 == 0 else WHITE
                para = cell.text_frame.paragraphs[0]
                para.font.size = Pt(11)
                para.font.color.rgb = DARK_TEXT
        self._footer(slide)
        return {"status": "ok", "slide": "table", "row_count": len(rows)}

    def add_recommendations_slide(self, heading: str, items: list[dict]):
        slide = self._new_slide()
        self._header(slide, heading, band_color=PRIMARY)
        box = slide.shapes.add_textbox(Inches(0.6), Inches(1.3), SLIDE_W - Inches(1.2), Inches(5.6))
        tf = box.text_frame
        tf.word_wrap = True
        priority_colors = {
            "high": RGBColor.from_string("B71C1C"),
            "medium": RGBColor.from_string("E65100"),
            "low": RGBColor.from_string("2E7D32"),
        }
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            priority = str(item.get("priority", "Medium"))
            p.text = f"[{priority.upper()}]  {item.get('action', '')}"
            p.font.size = Pt(15)
            p.font.bold = True
            p.font.color.rgb = priority_colors.get(priority.lower(), DARK_TEXT)
            p.space_before = Pt(10)
            p2 = tf.add_paragraph()
            p2.text = item.get("rationale", "")
            p2.font.size = Pt(12)
            p2.font.color.rgb = RGBColor.from_string("444444")
            p2.level = 1
        self._footer(slide)
        return {"status": "ok", "slide": "recommendations", "item_count": len(items)}

    def save(self, path):
        self.prs.save(path)


# Groq/OpenAI-format tool schemas exposed to Agent 2. Kept next to the
# implementation so the two can never drift apart.
TOOL_SCHEMAS = [
    {
        "type": "function",
        "function": {
            "name": "add_title_slide",
            "description": "Add the report's title slide. Call exactly once, first.",
            "parameters": {
                "type": "object",
                "properties": {
                    "title": {"type": "string"},
                    "subtitle": {"type": "string"},
                    "period": {"type": "string", "description": "Reporting period, e.g. 'July - August 2026'"},
                },
                "required": ["title", "subtitle", "period"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "add_kpi_slide",
            "description": "Add an executive-summary slide of KPI cards.",
            "parameters": {
                "type": "object",
                "properties": {
                    "heading": {"type": "string"},
                    "kpis": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "label": {"type": "string"},
                                "value": {"type": "string"},
                                "target": {"type": "string"},
                                "status": {"type": "string", "enum": ["on_target", "at_risk"]},
                            },
                            "required": ["label", "value", "target", "status"],
                        },
                    },
                },
                "required": ["heading", "kpis"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "add_narrative_slide",
            "description": "Add a slide of narrative paragraphs (performance overview or root-cause analysis).",
            "parameters": {
                "type": "object",
                "properties": {
                    "heading": {"type": "string"},
                    "paragraphs": {"type": "array", "items": {"type": "string"}},
                },
                "required": ["heading", "paragraphs"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "add_table_slide",
            "description": "Add a data table slide (e.g. top at-risk shipments, facility ranking). Use ONLY values passed to you in the DATA section -- do not invent rows.",
            "parameters": {
                "type": "object",
                "properties": {
                    "heading": {"type": "string"},
                    "columns": {"type": "array", "items": {"type": "string"}},
                    "rows": {"type": "array", "items": {"type": "array", "items": {"type": "string"}}},
                },
                "required": ["heading", "columns", "rows"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "add_recommendations_slide",
            "description": "Add the closing recommendations/action-items slide.",
            "parameters": {
                "type": "object",
                "properties": {
                    "heading": {"type": "string"},
                    "items": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "action": {"type": "string"},
                                "rationale": {"type": "string"},
                                "priority": {"type": "string", "enum": ["High", "Medium", "Low"]},
                            },
                            "required": ["action", "rationale", "priority"],
                        },
                    },
                },
                "required": ["heading", "items"],
            },
        },
    },
]
