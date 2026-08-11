"""
Raw-XML chart helpers for chart shapes python-pptx has no high-level builder
for: a combo chart (clustered bar on the primary axis + a line on a genuine
secondary value axis) and an XY-scatter with flat reference lines. Both are
built by first letting python-pptx create a normal single-series chart (so
the standard catAx/valAx/axId plumbing is correct), then splicing in the
extra series and axes as literal-cached XML -- no embedded-workbook range
references are needed for PowerPoint to render literal-cached data.

The secondary-axis pitfall this file exists to get right (see the pptx
authoring skill): a combo chart's second plot needs BOTH its own axId pair
declared -- a secondary valAx (visible, right side) AND a secondary catAx
(present but delete="1", hidden) -- not just a second value axis alone.
Skipping the hidden secondary catAx is exactly the mistake that makes
PowerPoint silently discard the chart as corrupt.
"""
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn

C_NS = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def _cat_str_lit(categories: list[str]) -> str:
    pts = "".join(f'<c:pt idx="{i}"><c:v>{_escape(c)}</c:v></c:pt>' for i, c in enumerate(categories))
    return f'<c:cat><c:strLit><c:ptCount val="{len(categories)}"/>{pts}</c:strLit></c:cat>'


def _val_num_lit(values: list[float]) -> str:
    pts = "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(values))
    return f'<c:val><c:numLit><c:formatCode>General</c:formatCode><c:ptCount val="{len(values)}"/>{pts}</c:numLit></c:val>'


def _escape(text: str) -> str:
    return (
        str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")
    )


def _next_axis_ids(plot_area) -> tuple[int, int]:
    existing = [int(el.get("val")) for el in plot_area.iter(qn("c:axId"))]
    base = max(existing) + 1 if existing else 500000000
    return base, base + 1


def add_combo_chart(
    slide,
    x,
    y,
    cx,
    cy,
    categories: list[str],
    bar_series_name: str,
    bar_values: list[float],
    line_series_name: str,
    line_values: list[float],
    bar_color_hex: str,
    line_color_hex: str,
):
    """Adds a clustered-bar + line combo chart (bar on the primary/left value
    axis, line on a real secondary/right value axis) to `slide`. Returns the
    GraphicFrame, same as `shapes.add_chart`."""
    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series(bar_series_name, bar_values)

    graphic_frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data)
    chart = graphic_frame.chart
    chart_space = chart._chartSpace
    plot_area = chart_space.find(f"{C_NS}chart").find(f"{C_NS}plotArea")

    bar_chart_el = plot_area.find(f"{C_NS}barChart")
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = RGBColor.from_string(bar_color_hex)

    primary_cat_id = int(plot_area.find(f"{C_NS}catAx").find(f"{C_NS}axId").get("val"))
    primary_val_id = int(plot_area.find(f"{C_NS}valAx").find(f"{C_NS}axId").get("val"))
    secondary_val_id, secondary_cat_id = _next_axis_ids(plot_area)

    line_chart_xml = f"""
    <c:lineChart xmlns:c="{C_NS[1:-1]}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
      <c:grouping val="standard"/>
      <c:varyColors val="0"/>
      <c:ser>
        <c:idx val="1"/>
        <c:order val="1"/>
        <c:tx><c:strRef><c:f></c:f><c:strCache><c:ptCount val="1"/><c:pt idx="0"><c:v>{_escape(line_series_name)}</c:v></c:pt></c:strCache></c:strRef></c:tx>
        <c:spPr><a:ln w="28575"><a:solidFill><a:srgbClr val="{line_color_hex}"/></a:solidFill></a:ln></c:spPr>
        <c:marker>
          <c:symbol val="circle"/>
          <c:size val="6"/>
          <c:spPr><a:solidFill><a:srgbClr val="{line_color_hex}"/></a:solidFill></c:spPr>
        </c:marker>
        {_cat_str_lit(categories)}
        {_val_num_lit(line_values)}
        <c:smooth val="0"/>
      </c:ser>
      <c:marker val="1"/>
      <c:axId val="{secondary_cat_id}"/>
      <c:axId val="{secondary_val_id}"/>
    </c:lineChart>
    """
    line_chart_el = parse_xml(line_chart_xml)
    bar_chart_el.addnext(line_chart_el)

    secondary_val_ax_xml = f"""
    <c:valAx xmlns:c="{C_NS[1:-1]}">
      <c:axId val="{secondary_val_id}"/>
      <c:scaling><c:orientation val="minMax"/></c:scaling>
      <c:delete val="0"/>
      <c:axPos val="r"/>
      <c:numFmt formatCode="General" sourceLinked="0"/>
      <c:majorTickMark val="out"/>
      <c:minorTickMark val="none"/>
      <c:tickLblPos val="nextTo"/>
      <c:txPr>
        <a:bodyPr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>
        <a:lstStyle xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>
        <a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:pPr><a:defRPr sz="900"/></a:pPr><a:endParaRPr lang="en-US"/></a:p>
      </c:txPr>
      <c:crossAx val="{secondary_cat_id}"/>
      <c:crosses val="max"/>
    </c:valAx>
    """
    secondary_cat_ax_xml = f"""
    <c:catAx xmlns:c="{C_NS[1:-1]}">
      <c:axId val="{secondary_cat_id}"/>
      <c:scaling><c:orientation val="minMax"/></c:scaling>
      <c:delete val="1"/>
      <c:axPos val="b"/>
      <c:majorTickMark val="out"/>
      <c:minorTickMark val="none"/>
      <c:tickLblPos val="nextTo"/>
      <c:crossAx val="{secondary_val_id}"/>
      <c:crosses val="autoZero"/>
      <c:auto val="1"/>
      <c:lblAlgn val="ctr"/>
      <c:lblOffset val="100"/>
      <c:noMultiLvlLbl val="0"/>
    </c:catAx>
    """
    plot_area.append(parse_xml(secondary_val_ax_xml))
    plot_area.append(parse_xml(secondary_cat_ax_xml))

    return graphic_frame
