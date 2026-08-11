"""
Builds a report deck on top of a user-uploaded PPTX template instead of
report_generator.ReportBuilder's own hand-positioned blank-layout shapes.
Exposes the SAME public method signatures as ReportBuilder (add_title_slide,
add_grouped_bar_slide, add_combo_slide, add_simple_chart_slide,
add_summary_slide, save_bytes) so report_generator.build_report can pick
whichever builder to use without the pivot-to-slide decision logic (in
_add_pivot_slides) needing to know or care which one it's talking to.

Content is placed into the TEMPLATE's own named layouts and their
placeholders -- title/subtitle on a "Cover" layout, a chart into a "Single
Chart" layout's content placeholder, narrative text into a "Single Column"
layout -- rather than fixed Inches() coordinates, since the template's own
slide size and branding (colors, fonts, footer, logo) aren't ours to assume.
Chart styling (data-label size, gridlines, axis titles) still goes through
report_generator's module-level style_* helpers so a template-based deck
looks consistent with the built-in one.

Falls back to picking the first available layout when a named one isn't
found -- an unfamiliar template still produces a deck, just a less
purpose-matched one, rather than failing outright.
"""
import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from app.schemas import OverallAnalysisReport
from app.services import chart_xml, report_generator as rg, report_style as style

COVER_LAYOUT_NAMES = ["Cover 1, example 1", "Cover 1", "Cover 2, example 1", "Cover 2"]
CHART_LAYOUT_NAMES = [f"Single Chart 0{i}" for i in range(1, 7)]
SUMMARY_LAYOUT_NAMES = ["Single Column 03", "Single Column 01", "Title Only"]

TITLE_IDX = 0
SUBTITLE_IDX = 1
# A chart layout's own body placeholders vary in idx per template; try the
# most title-adjacent ones first for a caption, most content-area one first
# for the chart's own bounding box.
CAPTION_IDX_CANDIDATES = (14, 13, 17, 12)
CONTENT_IDX_CANDIDATES = (1, 12, 13)
# The layout's own small side-label placeholder that sits inside the chart's
# now-enlarged bounding box (see _chart_bbox) -- left unfilled, it would show
# its own "click to add text" prompt right where the chart is drawn.
SIDE_LABEL_IDX_CANDIDATES = (13,)

CHART_MARGIN_X = Inches(0.4)
CHART_TOP_GAP = Inches(0.15)    # gap below the title
CHART_BOTTOM_GAP = Inches(0.1)  # gap above the caption
CAPTION_HEIGHT = Inches(0.4)      # enough for ~2 lines of small caption text
CAPTION_BOTTOM_MARGIN = Inches(0.3)  # gap above the layout's own slide-number footer


def _find_layout(master, names: list[str]):
    for name in names:
        for layout in master.slide_layouts:
            if layout.name == name:
                return layout
    return None


def _remove_existing_slides(prs) -> None:
    """Drops whatever example slides shipped inside the uploaded template --
    we only want ITS layouts/masters, not its own demo content."""
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.get(qn("r:id")))
        sldIdLst.remove(sldId)


def _placeholder(slide, *idx_candidates):
    by_idx = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
    for idx in idx_candidates:
        if idx in by_idx:
            return by_idx[idx]
    return None


def _remove_placeholder(slide, *idx_candidates) -> None:
    """Deletes a placeholder shape outright rather than leaving it unfilled.
    A slide auto-inherits every placeholder its layout has; one we have no
    content for (and aren't going to fill) would otherwise render its own
    "click to add text/content" prompt UNDER our chart -- since a native
    chart's own background is transparent outside its plot area, that prompt
    (and a content placeholder's four insert-icons) shows straight through."""
    ph = _placeholder(slide, *idx_candidates)
    if ph is not None:
        ph._element.getparent().remove(ph._element)


class TemplateReportBuilder:
    def __init__(self, template_bytes: bytes):
        self.prs = Presentation(io.BytesIO(template_bytes))
        master = self.prs.slide_masters[0]
        _remove_existing_slides(self.prs)

        self._cover_layout = _find_layout(master, COVER_LAYOUT_NAMES) or master.slide_layouts[0]
        chart_layouts = [_find_layout(master, [n]) for n in CHART_LAYOUT_NAMES]
        self._chart_layouts = [layout for layout in chart_layouts if layout is not None] or [master.slide_layouts[0]]
        self._summary_layout = _find_layout(master, SUMMARY_LAYOUT_NAMES) or master.slide_layouts[0]
        self._chart_layout_idx = 0
        self.slide_count = 0

    def _next_chart_layout(self):
        layout = self._chart_layouts[self._chart_layout_idx % len(self._chart_layouts)]
        self._chart_layout_idx += 1
        return layout

    def _new_slide(self, layout):
        slide = self.prs.slides.add_slide(layout)
        self.slide_count += 1
        return slide

    def _set_title(self, slide, text: str):
        title = _placeholder(slide, TITLE_IDX)
        if title is not None:
            title.text_frame.text = text

    def _set_caption(self, slide, text: str):
        """Moves the caption down near the slide's own footer -- freeing the
        space the layout originally left between the caption and the chart
        above it for a bigger chart (see _chart_bbox, which reads this new,
        lower position back out)."""
        caption = _placeholder(slide, *CAPTION_IDX_CANDIDATES)
        if caption is None:
            return
        # A placeholder with no explicit position of its own reports its
        # LAYOUT-inherited left/width when read -- but the moment any one
        # dimension is set, python-pptx materializes a fresh (zeroed) xfrm,
        # so left/width must be re-set too or they'd collapse to 0.
        left, width = caption.left, caption.width
        caption.top = self.prs.slide_height - CAPTION_BOTTOM_MARGIN - CAPTION_HEIGHT
        caption.height = CAPTION_HEIGHT
        caption.left = left
        caption.width = width
        tf = caption.text_frame
        tf.word_wrap = True
        tf.text = text

    def _chart_bbox(self, slide):
        """Returns a FULL-SIZE (left, top, width, height) box spanning
        nearly the whole slide width, and vertically from just below the
        title to just above the caption -- deliberately much larger than
        the layout's own (modest) content placeholder, per the requested
        "full size" chart. Removes that content placeholder (and the small
        side-label placeholder that would otherwise sit inside this
        enlarged area) since neither is used any more -- see
        _remove_placeholder's docstring for why leaving them unfilled would
        show through the chart."""
        title = _placeholder(slide, TITLE_IDX)
        top = (title.top + title.height + CHART_TOP_GAP) if title is not None else Inches(0.9)

        caption = _placeholder(slide, *CAPTION_IDX_CANDIDATES)
        bottom = (caption.top - CHART_BOTTOM_GAP) if caption is not None else (self.prs.slide_height - Inches(0.6))

        _remove_placeholder(slide, *CONTENT_IDX_CANDIDATES)
        _remove_placeholder(slide, *SIDE_LABEL_IDX_CANDIDATES)

        left = CHART_MARGIN_X
        width = self.prs.slide_width - 2 * CHART_MARGIN_X
        height = max(bottom - top, Inches(1))
        return left, top, width, height

    # -- slides, same public shape as report_generator.ReportBuilder --------

    def add_title_slide(self, title: str, subtitle: str | None):
        slide = self._new_slide(self._cover_layout)
        self._set_title(slide, title)
        if subtitle:
            sub = _placeholder(slide, SUBTITLE_IDX)
            if sub is not None:
                sub.text_frame.text = subtitle

    def add_grouped_bar_slide(self, heading: str, description: str, groups: list[tuple[str, list[tuple[str, float]]]],
                               y_axis_title: str, category_axis_title: str):
        slide = self._new_slide(self._next_chart_layout())
        self._set_title(slide, heading)
        self._set_caption(slide, description)
        left, top, width, height = self._chart_bbox(slide)

        data = CategoryChartData()
        values = []
        for level1_label, leaves in groups:
            category = data.categories.add_category(level1_label)
            for level2_label, value in leaves:
                category.add_sub_category(level2_label)
                values.append(value)
        data.add_series(y_axis_title, values)

        chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, left, top, width, height, data).chart
        rg.style_native_chart(chart, number_format="#,##0", single_series=True)
        rg.set_axis_title(chart.value_axis, y_axis_title)
        rg.set_axis_title(chart.category_axis, category_axis_title)

    def add_combo_slide(self, heading: str, description: str, categories: list[str], category_axis_title: str,
                         bar_name: str, bar_values: list[float], line_name: str, line_values: list[float]):
        slide = self._new_slide(self._next_chart_layout())
        self._set_title(slide, heading)
        self._set_caption(slide, description)
        left, top, width, height = self._chart_bbox(slide)

        chart = chart_xml.add_combo_chart(
            slide, left, top, width, height,
            categories=categories, bar_series_name=bar_name, bar_values=bar_values,
            line_series_name=line_name, line_values=line_values,
            bar_color_hex=style.BAR_COLOR_HEX, line_color_hex=style.LINE_COLOR_HEX,
            bar_axis_title=bar_name, line_axis_title=line_name,
        ).chart
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(style.CHART_FONT_PT)
        chart.legend.font.name = style.FONT_BODY
        for plot, fmt in zip(chart.plots, ('0"%"', "#,##0")):
            plot.has_data_labels = True
            plot.data_labels.number_format = fmt
            plot.data_labels.number_format_is_linked = False
            plot.data_labels.font.size = Pt(style.CHART_DATA_LABEL_FONT_PT)
            plot.data_labels.font.name = style.FONT_BODY
        rg.set_axis_title(chart.category_axis, category_axis_title)

    def add_simple_chart_slide(self, heading: str, description: str, categories: list[str], category_axis_title: str,
                                metric_label: str, values: list[float], is_trend: bool):
        slide = self._new_slide(self._next_chart_layout())
        self._set_title(slide, heading)
        self._set_caption(slide, description)
        left, top, width, height = self._chart_bbox(slide)

        data = CategoryChartData()
        data.categories = categories
        data.add_series(metric_label, values)

        chart_type = XL_CHART_TYPE.LINE_MARKERS if is_trend else XL_CHART_TYPE.COLUMN_CLUSTERED
        chart = slide.shapes.add_chart(chart_type, left, top, width, height, data).chart
        rg.style_native_chart(chart, number_format="#,##0.##", single_series=True)
        rg.set_axis_title(chart.value_axis, metric_label)
        rg.set_axis_title(chart.category_axis, category_axis_title)

    def add_summary_slide(self, heading: str, overall: OverallAnalysisReport | None):
        slide = self._new_slide(self._summary_layout)
        self._set_title(slide, heading)
        content = _placeholder(slide, *CONTENT_IDX_CANDIDATES)
        if content is None:
            return
        tf = content.text_frame
        tf.word_wrap = True
        if overall is None:
            tf.text = "No overall analysis had been generated for this session yet."
            return
        tf.text = overall.narrative
        for h in overall.highlights:
            p = tf.add_paragraph()
            p.text = f"•  {h.label}: {h.value}"

    def save_bytes(self) -> bytes:
        buffer = io.BytesIO()
        self.prs.save(buffer)
        return buffer.getvalue()
