"""The from-scratch (no user-uploaded template) .pptx builder: every slide,
heading, and chart is 100% rule-driven, hand-positioned onto a blank slide
layout. Used by report_generator.build_report only when neither a
user-uploaded template nor the bundled default template is available -- see
that module's docstring for the fallback order.
"""
import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.schemas import OverallAnalysisReport
from app.services import chart_xml, report_style as style
from app.services.report.chart_styling import (
    DARK_TEXT,
    MUTED,
    PRIMARY,
    WHITE,
    set_axis_title,
    style_native_chart,
)

TITLE_LAYOUT = 0
BLANK_LAYOUT = 6


class ReportBuilder:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = style.SLIDE_W
        self.prs.slide_height = style.SLIDE_H
        self._blank = self.prs.slide_layouts[BLANK_LAYOUT]
        self.slide_count = 0

    def _new_slide(self, bordered: bool = True):
        slide = self.prs.slides.add_slide(self._blank)
        self.slide_count += 1
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = WHITE
        self._side_band(slide)
        if bordered:
            self._add_border(slide)
        return slide

    def _add_border(self, slide):
        # A thin dark-navy frame just inside the slide edges -- every
        # content slide (Slide 2 onward) gets one; the cover slide doesn't.
        border = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, style.BORDER_MARGIN, style.BORDER_MARGIN,
            style.SLIDE_W - 2 * style.BORDER_MARGIN, style.SLIDE_H - 2 * style.BORDER_MARGIN,
        )
        border.fill.background()
        border.line.color.rgb = PRIMARY
        border.line.width = style.BORDER_WEIGHT
        border.shadow.inherit = False

    def _side_band(self, slide):
        # Solid navy strip down the slide's right edge, matching the Carrier
        # template's own master slide -- every content box already keeps a
        # 0.5in right margin, well clear of this band.
        band = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, style.SLIDE_W - style.SIDE_BAND_WIDTH, 0, style.SIDE_BAND_WIDTH, style.SLIDE_H
        )
        band.fill.solid()
        band.fill.fore_color.rgb = PRIMARY
        band.line.fill.background()
        band.shadow.inherit = False

    def _header(self, slide, heading: str):
        # Plain left-aligned text on the white slide background -- no
        # full-width color band. A decorative bar spanning the slide reads
        # as filler, and the reference deck's own clean look skips it too.
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.22), style.SLIDE_W - Inches(1.0), style.HEADER_HEIGHT)
        tf = box.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = heading
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        p.font.name = style.FONT_TITLE

    def _footer(self, slide):
        logo_top = style.SLIDE_H - Inches(0.365)
        slide.shapes.add_picture(str(style.LOGO_PATH), Inches(0.4), logo_top, style.LOGO_WIDTH, style.LOGO_HEIGHT)

        tb = slide.shapes.add_textbox(Inches(1.25), style.SLIDE_H - Inches(0.35), Inches(2.5), Inches(0.25))
        p = tb.text_frame.paragraphs[0]
        p.text = style.PROGRAM_TITLE
        p.font.size = Pt(12)
        p.font.color.rgb = MUTED
        p.font.name = style.FONT_BODY

        proprietary = slide.shapes.add_textbox(Inches(3.85), style.SLIDE_H - Inches(0.35), Inches(2.3), Inches(0.25))
        pr = proprietary.text_frame.paragraphs[0]
        pr.text = style.PROPRIETARY_TEXT
        pr.font.size = Pt(12)
        pr.alignment = PP_ALIGN.CENTER
        pr.font.color.rgb = MUTED
        pr.font.name = style.FONT_BODY

        page = slide.shapes.add_textbox(Inches(8.6), style.SLIDE_H - Inches(0.35), Inches(0.9), Inches(0.25))
        pp = page.text_frame.paragraphs[0]
        pp.text = str(self.slide_count)
        pp.font.size = Pt(12)
        pp.alignment = PP_ALIGN.RIGHT
        pp.font.color.rgb = MUTED
        pp.font.name = style.FONT_BODY

    def _caption(self, slide, text: str, top):
        box = slide.shapes.add_textbox(Inches(0.5), top, style.SLIDE_W - Inches(1.0), Inches(0.35))
        tf = box.text_frame
        tf.word_wrap = True
        tf.text = text
        run = tf.paragraphs[0].runs[0]
        run.font.size = Pt(12)
        run.font.color.rgb = MUTED
        run.font.name = style.FONT_BODY

    # -- slides ---------------------------------------------------------------
    def add_title_slide(self, title: str, subtitle: str | None):
        """Cover slide: logo top-left, title/subtitle on the left, and a
        photo-collage-shaped block of flat colour on the right -- no actual
        photography is available for this deck, so solid colour panels stand
        in for where a Carrier-branded cover would place its imagery, rather
        than faking or downloading photos."""
        slide = self._new_slide(bordered=False)

        slide.shapes.add_picture(str(style.LOGO_PATH), Inches(0.4), Inches(0.35), style.TITLE_LOGO_WIDTH, style.TITLE_LOGO_HEIGHT)

        tf = slide.shapes.add_textbox(Inches(0.4), Inches(3.1), Inches(3.9), Inches(1.6)).text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        p.font.name = style.FONT_TITLE
        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.size = Pt(13)
            p2.font.color.rgb = MUTED
            p2.font.name = style.FONT_BODY

        # Photo-collage stand-in: one large panel above two smaller ones,
        # occupying the right half of the slide (clear of the side band).
        collage_x, collage_w = Inches(4.6), style.SLIDE_W - Inches(4.6) - Inches(0.3)
        large = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, collage_x, Inches(0.5), collage_w, Inches(3.6))
        large.fill.solid()
        large.fill.fore_color.rgb = PRIMARY
        large.line.fill.background()
        large.shadow.inherit = False

        gap = Inches(0.2)
        small_w = (collage_w - gap) / 2
        for i, hex_color in enumerate((style.BAR_PALETTE_HEX[3], style.BAR_PALETTE_HEX[4])):  # teal, amber
            small = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, collage_x + i * (small_w + gap), Inches(4.3), small_w, Inches(2.4)
            )
            small.fill.solid()
            small.fill.fore_color.rgb = RGBColor.from_string(hex_color)
            small.line.fill.background()
            small.shadow.inherit = False

    def add_grouped_bar_slide(self, heading: str, description: str, groups: list[tuple[str, list[tuple[str, float]]]],
                               y_axis_title: str, category_axis_title: str):
        """`groups` is level1 -> [(level2, value), ...], as built by
        `_grouped_bar_data`. Rendered as a single-series, multi-level
        category-axis chart -- one navy bar per (level1, level2) pair,
        grouped visually by level1 -- matching an Excel PivotChart's own
        two-level axis rather than a colour-coded legend."""
        slide = self._new_slide()
        self._header(slide, heading)
        self._caption(slide, description, top=Inches(0.6))

        data = CategoryChartData()
        values = []
        for level1_label, leaves in groups:
            category = data.categories.add_category(level1_label)
            for level2_label, value in leaves:
                category.add_sub_category(level2_label)
                values.append(value)
        data.add_series(y_axis_title, values)

        gf = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), style.CHART_TOP,
            style.SLIDE_W - Inches(1.0), style.SLIDE_H - style.CHART_TOP - Inches(0.6), data,
        )
        chart = gf.chart
        style_native_chart(chart, number_format="#,##0", single_series=True)
        set_axis_title(chart.value_axis, y_axis_title)
        set_axis_title(chart.category_axis, category_axis_title)
        self._footer(slide)

    def add_combo_slide(self, heading: str, description: str, categories: list[str], category_axis_title: str,
                         bar_name: str, bar_values: list[float], line_name: str, line_values: list[float]):
        slide = self._new_slide()
        self._header(slide, heading)
        self._caption(slide, description, top=Inches(0.6))

        gf = chart_xml.add_combo_chart(
            slide, Inches(0.5), style.CHART_TOP, style.SLIDE_W - Inches(1.0), style.SLIDE_H - style.CHART_TOP - Inches(0.6),
            categories=categories, bar_series_name=bar_name, bar_values=bar_values,
            line_series_name=line_name, line_values=line_values,
            bar_color_hex=style.BAR_COLOR_HEX, line_color_hex=style.LINE_COLOR_HEX,
            bar_axis_title=bar_name, line_axis_title=line_name,
        )
        chart = gf.chart
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(style.CHART_FONT_PT)
        chart.legend.font.name = style.FONT_BODY
        for plot, fmt in zip(chart.plots, ('0"%"', "#,##0")):
            plot.has_data_labels = True
            plot.data_labels.number_format = fmt
            plot.data_labels.number_format_is_linked = False
            plot.data_labels.font.size = Pt(style.CHART_DATA_LABEL_FONT_PT)
            plot.data_labels.font.name = style.FONT_BODY
        # category_axis/tick fonts and the primary value axis's title,
        # gridlines, axis line, and tick font are ALL set inside
        # add_combo_chart -- once a secondary c:valAx exists, python-pptx's
        # own chart.value_axis resolves to the secondary one here, not the
        # primary, so touching axis styling from this side would silently
        # land on the wrong axis (see add_combo_chart's docstring).
        set_axis_title(chart.category_axis, category_axis_title)
        self._footer(slide)

    def add_simple_chart_slide(self, heading: str, description: str, categories: list[str], category_axis_title: str,
                                metric_label: str, values: list[float], is_trend: bool):
        slide = self._new_slide()
        self._header(slide, heading)
        self._caption(slide, description, top=Inches(0.6))

        data = CategoryChartData()
        data.categories = categories
        data.add_series(metric_label, values)

        chart_type = XL_CHART_TYPE.LINE_MARKERS if is_trend else XL_CHART_TYPE.COLUMN_CLUSTERED
        gf = slide.shapes.add_chart(
            chart_type, Inches(0.5), style.CHART_TOP, style.SLIDE_W - Inches(1.0), style.SLIDE_H - style.CHART_TOP - Inches(0.6), data,
        )
        chart = gf.chart
        style_native_chart(chart, number_format="#,##0.##", single_series=True)
        set_axis_title(chart.value_axis, metric_label)
        set_axis_title(chart.category_axis, category_axis_title)
        self._footer(slide)

    def add_summary_slide(self, heading: str, overall: OverallAnalysisReport | None):
        slide = self._new_slide()
        self._header(slide, heading)

        if overall is None:
            self._caption(slide, "No overall analysis had been generated for this session yet.", top=Inches(0.6))
            self._footer(slide)
            return

        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.75), style.SLIDE_W - Inches(1.0), Inches(1.5))
        tf = box.text_frame
        tf.word_wrap = True
        tf.text = overall.narrative
        tf.paragraphs[0].font.size = Pt(12)
        tf.paragraphs[0].font.color.rgb = DARK_TEXT
        tf.paragraphs[0].font.name = style.FONT_BODY

        list_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.4), style.SLIDE_W - Inches(1.0), style.SLIDE_H - Inches(2.9))
        tf2 = list_box.text_frame
        tf2.word_wrap = True
        for i, h in enumerate(overall.highlights):
            p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
            p.text = f"•  {h.label}: {h.value}"
            p.font.size = Pt(12)
            p.font.color.rgb = DARK_TEXT
            p.font.name = style.FONT_BODY
            p.space_after = Pt(8)
        self._footer(slide)

    def save_bytes(self) -> bytes:
        buffer = io.BytesIO()
        self.prs.save(buffer)
        return buffer.getvalue()
