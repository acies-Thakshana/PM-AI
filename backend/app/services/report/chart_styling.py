"""Pure/near-pure helpers for shaping pivot rows into chart data and styling
native python-pptx chart objects -- kept separate from slide-building and
slide-shape-selection so a chart-styling change never risks touching either
of those. Module-level (not a class) so report_template_builder.py can
reuse the same styling for a template-based deck.
"""
import re
from collections import OrderedDict

import pandas as pd
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.util import Pt

from app.schemas import PivotResult
from app.services import report_style as style

MAX_CHART_ROWS = 20

PRIMARY = RGBColor.from_string(style.BAR_COLOR_HEX)
LINE_COLOR = RGBColor.from_string(style.LINE_COLOR_HEX)
WHITE = RGBColor.from_string(style.WHITE_HEX)
DARK_TEXT = RGBColor.from_string(style.DARK_TEXT_HEX)
MUTED = RGBColor.from_string(style.MUTED_TEXT_HEX)
BAR_PALETTE = [RGBColor.from_string(h) for h in style.BAR_PALETTE_HEX]

_ACTIVITY_DATE_HINTS = re.compile(r"arriv|depart|segment|trip create", re.IGNORECASE)


def _date_range_label(df: pd.DataFrame) -> str | None:
    date_cols = list(df.select_dtypes(include=["datetime64[ns]", "datetimetz"]).columns)
    if not date_cols:
        return None
    activity_cols = [c for c in date_cols if _ACTIVITY_DATE_HINTS.search(str(c))]
    date_cols = activity_cols or date_cols

    overall_min, overall_max = None, None
    for col in date_cols:
        col_min, col_max = df[col].min(), df[col].max()
        if pd.isna(col_min) or pd.isna(col_max):
            continue
        overall_min = col_min if overall_min is None or col_min < overall_min else overall_min
        overall_max = col_max if overall_max is None or col_max > overall_max else overall_max
    if overall_min is None or overall_max is None:
        return None
    return f"{overall_min:%d.%m.%Y} – {overall_max:%d.%m.%Y}"


def _is_trend_pivot(pivot: PivotResult) -> bool:
    return len(pivot.group_by) == 1 and "month" in pivot.group_by[0].lower()


def _numeric_rows(rows: list[dict], metric_label: str) -> list[dict]:
    return [r for r in rows if isinstance(r.get(metric_label), (int, float))]


def _blank_label(value) -> str:
    """Mirrors Excel's own PivotChart convention: a missing group value reads
    as "(blank)" rather than the Python str() of None/NaN leaking through."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return style.BLANK_LABEL
    text = str(value).strip()
    if not text or text.lower() in ("none", "nan"):
        return style.BLANK_LABEL
    return text


def _grouped_bar_data(rows: list[dict], level1_col: str, level2_col: str, metric_label: str):
    """Shapes two-level pivot rows into level1 -> [(level2, value), ...]
    groups for a multi-level-category bar chart: one bar per (level1, level2)
    pair, grouped on the category axis by level1 -- this is an Excel
    PivotChart's own two-level axis, not a colour-coded legend, so every
    individual level2 value (e.g. grower) gets its own labelled bar instead
    of being folded into an "Other" series.

    Groups sort alphabetically with the blank/missing group pushed last;
    leaves within a group sort by value descending -- matches the reference
    deck's own ordering. Total bar count is capped at MULTI_LEVEL_MAX_LEAVES,
    dropping the lowest-value pairs first."""
    groups: "OrderedDict[str, list[tuple[str, float]]]" = OrderedDict()
    for r in _numeric_rows(rows, metric_label):
        l1, l2 = _blank_label(r.get(level1_col)), _blank_label(r.get(level2_col))
        groups.setdefault(l1, []).append((l2, r[metric_label]))

    ordered_keys = sorted(groups, key=lambda k: (k == style.BLANK_LABEL, k.lower()))
    result = [(key, sorted(groups[key], key=lambda pair: -pair[1])) for key in ordered_keys]

    leaf_total = sum(len(leaves) for _, leaves in result)
    if leaf_total > style.MULTI_LEVEL_MAX_LEAVES:
        flat = [(key, l2, v) for key, leaves in result for l2, v in leaves]
        kept = {(key, l2) for key, l2, v in sorted(flat, key=lambda t: -t[2])[: style.MULTI_LEVEL_MAX_LEAVES]}
        result = [(key, [(l2, v) for l2, v in leaves if (key, l2) in kept]) for key, leaves in result]
        result = [(key, leaves) for key, leaves in result if leaves]

    return result


def _combo_data(rows: list[dict], group_col: str, pct_label: str, count_label: str, sort_by_count: bool):
    scored = [r for r in rows if isinstance(r.get(pct_label), (int, float)) and isinstance(r.get(count_label), (int, float))]
    # A ranking pivot (e.g. carriers) reads best ordered by volume; a genuine
    # time trend (e.g. by month) would have that chronological order
    # scrambled by the same sort, so it keeps whatever order pivot_engine's
    # own sort_by already produced (chronological, per the Analysis Profile).
    if sort_by_count:
        scored.sort(key=lambda r: -r[count_label])
    scored = scored[:MAX_CHART_ROWS]
    categories = [str(r.get(group_col)) for r in scored]
    return categories, [r[pct_label] for r in scored], [r[count_label] for r in scored]


def style_axes_grid(category_axis, value_axis):
    """Left value-axis line + bottom category-axis line, both made visible,
    plus horizontal (value) and vertical (category) major gridlines --
    frames the plot without a full boxed border."""
    grid_color = RGBColor.from_string(style.GRIDLINE_COLOR_HEX)
    for axis in (category_axis, value_axis):
        axis.has_major_gridlines = True
        axis.has_minor_gridlines = False
        axis.major_gridlines.format.line.color.rgb = grid_color
        axis.major_gridlines.format.line.width = Pt(0.75)
        axis.format.line.color.rgb = MUTED
        axis.format.line.width = Pt(1)


def set_axis_title(axis, text: str):
    axis.axis_title.text_frame.text = text
    axis.axis_title.text_frame.paragraphs[0].font.size = Pt(style.CHART_AXIS_TITLE_FONT_PT)
    axis.axis_title.text_frame.paragraphs[0].font.name = style.FONT_BODY


def style_native_chart(chart, number_format: str, single_series: bool):
    style.set_chart_default_font(chart, style.CHART_DATA_LABEL_FONT_PT, style.FONT_BODY)
    chart.has_title = False
    chart.has_legend = not single_series
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(style.CHART_FONT_PT)
        chart.legend.font.name = style.FONT_BODY
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format = number_format
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.size = Pt(style.CHART_DATA_LABEL_FONT_PT)
    plot.data_labels.font.name = style.FONT_BODY
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    for i, series in enumerate(plot.series):
        color = BAR_PALETTE[i % len(BAR_PALETTE)]
        if chart.chart_type == XL_CHART_TYPE.LINE_MARKERS:
            series.format.line.color.rgb = color
            series.format.line.width = Pt(2.25)
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = color
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color
    style_axes_grid(chart.category_axis, chart.value_axis)
    chart.category_axis.tick_labels.font.size = Pt(style.CHART_FONT_PT)
    chart.category_axis.tick_labels.font.name = style.FONT_BODY
    chart.value_axis.tick_labels.font.size = Pt(style.CHART_FONT_PT)
    chart.value_axis.tick_labels.font.name = style.FONT_BODY
